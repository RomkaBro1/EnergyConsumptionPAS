"""Build reviewable DOCX/PDF deliverables from versioned source documents."""
import re
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from playwright.sync_api import sync_playwright

base = Path(__file__).resolve().parents[1]
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin, sec.right_margin = Cm(2.5), Cm(2)
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(7)
style.paragraph_format.line_spacing = 1.15
for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
    doc.styles[name].font.name = "Calibri"
    doc.styles[name].font.color.rgb = RGBColor.from_string("176D67")

lines = (base / "docs/report.md").read_text(encoding="utf-8").splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1
        continue
    if line.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.fullmatch(r"[|:\-\s]+", lines[i]):
                rows.append([c.strip().replace("`", "") for c in lines[i].strip("|").split("|")])
            i += 1
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = "Light Shading Accent 1"
        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = value
        continue
    if line.startswith("#"):
        level = len(line) - len(line.lstrip("#"))
        doc.add_heading(line.lstrip("# "), 0 if level == 1 else level - 1)
    else:
        doc.add_paragraph(line.replace("`", "").replace("**", ""))
    i += 1
doc.core_properties.title = "GridPulse DE — прогнозно-аналитическая система"
doc.core_properties.subject = "Проект по дисциплине Прогнозно-аналитические системы"
doc.save(base / "docs/report.docx")

with sync_playwright() as p:
    browser = p.chromium.launch(channel="msedge", headless=True)
    page = browser.new_page(viewport={"width": 1280, "height": 720})
    page.goto((base / "docs/presentation.html").as_uri(), wait_until="networkidle")
    page.pdf(path=str(base / "docs/presentation.pdf"), width="1280px", height="720px", print_background=True, prefer_css_page_size=True)
    browser.close()
print("Created docs/report.docx and docs/presentation.pdf")
