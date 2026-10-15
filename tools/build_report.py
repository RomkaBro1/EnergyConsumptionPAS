"""Build the coursework report from Markdown and the author's original cover.

Run this script, then tools/render_report.ps1 -Update to fill the Word TOC and PDF.
The original source document is never saved or changed.
"""
import json
import re
from copy import deepcopy
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parents[1]
ASSETS = BASE / 'docs/report-assets'
ASSETS.mkdir(exist_ok=True)
plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 12})


def diagrams():
    fig, ax = plt.subplots(figsize=(12, 4.5))
    ax.set(xlim=(0, 12), ylim=(0, 4.5)); ax.axis('off')
    boxes = [(0.1, 2.6, 2.2, 1.4, 'Источники\nEnergy-Charts\nDWD · Feiertage'),
             (3, 2.6, 2.2, 1.4, 'Bronze\nИсходные ответы\nURL · время · SHA-256'),
             (6, 2.6, 2.2, 1.4, 'Silver · SQLite\nЗаписи · версии\nЗапуски · проверки'),
             (9, 2.6, 2.7, 1.4, 'Gold\nЧасовая витрина\nКалендарь · признаки'),
             (9, .25, 2.7, 1.4, 'Две модели\nНедельный профиль\nГрадиентный бустинг'),
             (5.1, .25, 2.8, 1.4, 'Flask API\nПрогноз · сценарии\nCSV · происхождение'),
             (.1, .25, 3.6, 1.4, 'Интерфейс\nПредметный дашборд\nОперационный дашборд')]
    for x,y,w,h,text in boxes:
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.06',fc='#f2f4f6',ec='#343b43',lw=1.2))
        ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=11)
    for start,end in [((2.4,3.3),(2.9,3.3)),((5.3,3.3),(5.9,3.3)),((8.3,3.3),(8.9,3.3)),((10.35,2.5),(10.35,1.75)),((8.9,.95),(8,.95)),((5,.95),(3.8,.95))]:
        ax.add_patch(FancyArrowPatch(start,end,arrowstyle='-|>',mutation_scale=16,color='#343b43'))
    ax.text(5.8,2.1,'Загрузчики и парсеры → контроль → объединение по UTC',ha='center',fontsize=10)
    fig.savefig(ASSETS/'architecture.png',dpi=220,bbox_inches='tight',facecolor='white');plt.close(fig)
    fig,ax=plt.subplots(figsize=(12,2.6));ax.axis('off');ax.set(xlim=(0,12),ylim=(0,2.6))
    for x,w,color,title,detail in [(0,5.7,'#dce5ef','Обучение проверяемой модели','до 10.09.2026 16:00 UTC'),(5.9,2.8,'#e7e7e7','Калибровка · 336 ч','10.09 17:00 — 24.09 16:00'),(8.9,3.1,'#e0e9df','Тест · 336 ч','24.09 17:00 — 08.10 16:00')]:
        ax.add_patch(FancyBboxPatch((x,.8),w,1.3,boxstyle='round,pad=.02',fc=color,ec='#444'))
        ax.text(x+w/2,1.63,title,ha='center',fontsize=10,fontweight='bold');ax.text(x+w/2,1.15,detail,ha='center',fontsize=9)
    ax.text(6,.3,'Недельные старты · только предшествующие факты и допустимые выпуски погоды',ha='center',fontsize=11)
    fig.savefig(ASSETS/'validation-scheme.png',dpi=220,bbox_inches='tight',facecolor='white');plt.close(fig)


def slides():
    items = [
      ('Цель и результат', ['GridPulse DE: локальная система для нагрузки Германии', 'Мониторинг, прогноз на 24 / 72 / 168 часов', 'Погода и производственный календарь', 'Пользователь: аналитик; сопровождение: оператор данных']),
      ('Жизненный цикл данных', ['Energy-Charts: Load · DWD: наблюдения и прогнозы', 'Feiertage API и python-holidays: календарь', 'Bronze → Silver → Gold → модели → дашборды', 'Инкрементальность, идемпотентность, журнал, история']),
      ('Модели и проверка', ['Базовая модель: нагрузка 168 часов назад', 'Бустинг: 16 признаков, 180 итераций', '28 дней: 14 на калибровку и 14 на независимый тест', 'Погода: только выпуски, доступные на старте прогноза']),
      ('Результаты основного теста', ['Тест: 24 сентября — 8 октября 2026 года, 336 часов', 'MAE: бустинг 1624,23 МВт; профиль 1654,81 МВт', 'RMSE: бустинг 2136,01 МВт; профиль 2084,50 МВт', 'Покрытие: 87,50% и 89,88%; преимущество зависит от метрики']),
      ('Работа пользователя', ['Обзор: ожидаемый пик, история, энергия', 'Прогноз: три погодных параметра и выгрузка CSV', 'Качество: сравнение ошибок и условий эксперимента', 'Операционный раздел: свежесть, проверки, происхождение']),
      ('Проверки и ограничения', ['30 автоматических тестов и браузерная проверка', 'Крайние сценарии, CSV, мобильная ширина 390 пикселей', 'Равные веса городов; неполный архив погодных выпусков', 'Дальнейшая работа: сезоны, задержки публикации, мониторинг'])]
    for i,(title,lines) in enumerate(items,1):
        fig,ax=plt.subplots(figsize=(12.8,7.2));fig.patch.set_facecolor('white');ax.axis('off')
        ax.text(.05,.88,title,transform=ax.transAxes,fontsize=25,fontweight='bold',color='#222')
        ax.plot([.05,.96],[.79,.79],transform=ax.transAxes,color='#444',lw=1)
        for j,line in enumerate(lines):ax.text(.06,.66-j*.145,'•  '+line,transform=ax.transAxes,fontsize=16,va='top')
        ax.text(.05,.045,'Янушкевич Р.В. · ИНБО-20-23 · ПАС · 2026',transform=ax.transAxes,fontsize=11,color='#555')
        ax.text(.96,.045,str(i),ha='right',transform=ax.transAxes,fontsize=11)
        fig.savefig(ASSETS/f'slide-{i}.png',dpi=150,facecolor='white');plt.close(fig)


def style(doc,name,size=14,bold=False,align=WD_ALIGN_PARAGRAPH.JUSTIFY,first=1.25,before=0,after=0,line=1.5):
    s=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
    for attr in ('ascii','hAnsi','eastAsia','cs'):s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:'+attr),'Times New Roman')
    p=s.paragraph_format;p.alignment=align;p.first_line_indent=Cm(first);p.left_indent=Cm(0);p.right_indent=Cm(0)
    p.space_before=Pt(before);p.space_after=Pt(after);p.line_spacing=line;p.widow_control=True
    return s


def field(paragraph,instruction,text=''):
    r=paragraph.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r._r.append(begin)
    t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' '+instruction+' ';r._r.append(t)
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
    paragraph.add_run(text)
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');paragraph.add_run()._r.append(end)


def typography(text):
    # Nonbreaking spaces bind values and units; paths and URLs remain untouched.
    return re.sub(r'(?<=\d) (?=(?:ГВт|МВт|Вт|м/с|°C|час|мин|пиксел|дн)[\w·/²-]*)','\u00a0',text)


def main():
    diagrams();slides()
    source=BASE/'ПАС_ЯнушкевичРВ.docx'
    doc=Document(source)
    body=doc.element.body
    # Keep the original cover XML, tables, emblem, styles and package relationships.
    # The original contents begins at body child 22.
    children=list(body)
    for child in children[22:]:
        if child.tag!=qn('w:sectPr'):body.remove(child)
    for p in doc.paragraphs:
        for node in p._p.xpath('./w:pPr/w:sectPr'):node.getparent().remove(node)
    section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(2);section.left_margin=Cm(3);section.right_margin=Cm(1)
    # Clear legacy numbering; cover and contents count but do not display a number.
    for node in section._sectPr.findall(qn('w:pgNumType')):section._sectPr.remove(node)
    section.different_first_page_header_footer=False
    for area in (section.header,section.footer):
        for e in list(area._element):area._element.remove(e)
        area.add_paragraph()
    style(doc,'GP Body')
    for level,size in [(1,18),(2,16),(3,14)]:
        s=style(doc,f'GP Heading {level}',size,True,WD_ALIGN_PARAGRAPH.LEFT,0,0 if level==1 else 42.55,28.35)
        s.paragraph_format.left_indent=Cm(1.25);s.paragraph_format.keep_with_next=True;s.paragraph_format.keep_together=True
        s.paragraph_format.page_break_before=(level==1)
        outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),str(level-1));s.element.get_or_add_pPr().append(outline)
    style(doc,'GP Caption',12,True,WD_ALIGN_PARAGRAPH.CENTER,0,0,17,1)
    tablecap=style(doc,'GP Table Caption',12,False,WD_ALIGN_PARAGRAPH.LEFT,0,17,0,1);tablecap.font.italic=True;tablecap.paragraph_format.keep_with_next=True
    style(doc,'GP Table',12,False,WD_ALIGN_PARAGRAPH.LEFT,0,0,0,1)
    style(doc,'GP Figure',12,False,WD_ALIGN_PARAGRAPH.CENTER,0,0,0,1).paragraph_format.keep_with_next=True
    style(doc,'GP Bibliography',14,False,WD_ALIGN_PARAGRAPH.JUSTIFY,0,0,0,1.5)
    style(doc,'GP TOC',14,False,WD_ALIGN_PARAGRAPH.LEFT,0,0,0,1.5)
    style(doc,'GP Formula',14,False,WD_ALIGN_PARAGRAPH.CENTER,0,0,0,1.5)
    style(doc,'GP Where',14,False,WD_ALIGN_PARAGRAPH.LEFT,0,0,0,1.5)
    code=style(doc,'GP Code',10,False,WD_ALIGN_PARAGRAPH.LEFT,0,0,0,1);code.font.name='Courier New'
    for i in range(1,4):
        name=f'TOC {i}'
        if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(14);s.font.color.rgb=RGBColor(0,0,0)
        pf=s.paragraph_format;pf.first_line_indent=Cm(0);pf.left_indent=Cm(0);pf.space_after=Pt(0);pf.space_before=Pt(0);pf.line_spacing=1.5
    p=doc.add_paragraph('СОДЕРЖАНИЕ','GP Heading 1');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    # Keep the contents title outside the TOC itself.
    outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),'9');p._p.get_or_add_pPr().append(outline)
    p=doc.add_paragraph(style='GP TOC');field(p,'TOC \\o "1-3" \\h \\z \\u','Содержание обновляется средствами Microsoft Word')
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.footer.is_linked_to_previous=False;sec.header.is_linked_to_previous=False
    for node in sec._sectPr.findall(qn('w:pgNumType')):sec._sectPr.remove(node)
    for e in list(sec.footer._element):sec.footer._element.remove(e)
    fp=sec.footer.add_paragraph();fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.paragraph_format.first_line_indent=Cm(0)
    field(fp,'PAGE')
    for r in fp.runs:r.font.name='Times New Roman';r.font.size=Pt(11)
    lines=(BASE/'docs/report.md').read_text(encoding='utf-8').splitlines()
    i=0;first=True;after_table=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('# '):
            p=doc.add_paragraph(line[2:],'GP Heading 1')
            if first:p.paragraph_format.page_break_before=False;first=False
        elif line.startswith('## '):doc.add_paragraph(line[3:],'GP Heading 2')
        elif line.startswith('@table '):
            doc.add_paragraph(line[7:],'GP Table Caption');rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                if not re.fullmatch(r'[|:\-\s]+',lines[i]):rows.append([v.strip() for v in lines[i].strip('|').split('|')])
                i+=1
            table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid';table.autofit=True
            for ri,row in enumerate(rows):
                cells=table.add_row().cells
                for cell,value in zip(cells,row):
                    p=cell.paragraphs[0];p.style='GP Table';r=p.add_run(typography(value));r.bold=(ri==0)
                pr=table.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
                if ri==0:pr.append(OxmlElement('w:tblHeader'))
                # Tables are short enough to remain whole; no unlabelled continuation.
                for cell in cells:
                    for p in cell.paragraphs:p.paragraph_format.keep_with_next=(ri<len(rows)-1)
            after_table=True
        elif line.startswith('@figure '):
            path,caption=line[8:].split(' | ',1);p=doc.add_paragraph(style='GP Figure')
            im=Image.open(BASE/'docs'/path);w,h=im.size;maxh=19 if 'scenario' in path else 14
            width=min(17,maxh*w/h);p.add_run().add_picture(str(BASE/'docs'/path),width=Cm(width))
            doc.add_paragraph(caption,'GP Caption')
        elif line.startswith('@formula '):
            number,expr=line[9:].split(' | ',1)
            doc.add_paragraph('',style='GP Body').paragraph_format.line_spacing=1
            p=doc.add_paragraph(style='GP Formula');p.paragraph_format.tab_stops.add_tab_stop(Cm(8.5),WD_TAB_ALIGNMENT.CENTER);p.paragraph_format.tab_stops.add_tab_stop(Cm(17),WD_TAB_ALIGNMENT.RIGHT)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.add_run('\t')
            math=OxmlElement('m:oMath');mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text=expr;mr.append(mt);math.append(mr);p._p.append(math)
            p.add_run('\t('+number+')')
            doc.add_paragraph('',style='GP Body').paragraph_format.line_spacing=1
        elif line.startswith('@where '):
            p=doc.add_paragraph(typography(line[7:].replace('\\t','\t').replace('\\n','\n')),'GP Where')
        elif line.startswith('@bibgroup '):
            p=doc.add_paragraph(line[10:].upper(),'GP Where');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.left_indent=Cm(1.25);p.paragraph_format.space_before=Pt(17);p.paragraph_format.space_after=Pt(17);p.paragraph_format.keep_with_next=True
        elif line.startswith('@ref '):
            number,text=line[5:].split(' ',1);p=doc.add_paragraph(number+'\t'+text,'GP Bibliography');p.paragraph_format.tab_stops.add_tab_stop(Cm(1.25));p.paragraph_format.left_indent=Cm(1.25);p.paragraph_format.first_line_indent=Cm(-1.25)
        elif line.startswith('@appendix '):
            letter,title=line[10:].split(' | ',1);p=doc.add_paragraph('Приложение '+letter,'GP Heading 3');p.paragraph_format.page_break_before=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.left_indent=Cm(0)
            p=doc.add_paragraph(title,'GP Body');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
        elif line.startswith('@code '):
            doc.add_paragraph(line[6:],'GP Table Caption');content=[]
            while i<len(lines) and lines[i]!='@endcode':content.append(lines[i]);i+=1
            i+=1;table=doc.add_table(rows=1,cols=1);table.style='Table Grid';p=table.cell(0,0).paragraphs[0];p.style='GP Code';p.add_run('\n'.join(content));after_table=True
        else:
            p=doc.add_paragraph(typography(line),'GP Body')
            if after_table:p.paragraph_format.space_before=Pt(17);after_table=False
    doc.core_properties.title='Прогнозно-аналитическая система мониторинга почасового потребления электроэнергии Германии'
    doc.core_properties.author='Янушкевич Р.В.'
    doc.core_properties.subject='Проектная работа по дисциплине «Прогнозно-аналитические системы»'
    doc.core_properties.keywords='GridPulse DE; ПАС; Германия; нагрузка; прогнозирование'
    doc.save(BASE/'docs/report.docx')
    print('Built docs/report.docx; update fields and export with tools/render_report.ps1')


if __name__=='__main__':main()
